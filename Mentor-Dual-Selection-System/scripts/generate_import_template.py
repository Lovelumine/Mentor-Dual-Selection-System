#!/usr/bin/env python3
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TEMPLATES = {
    "teacher": {
        "sheet_title": "导师导入模板",
        "filename": "导师信息导入模板.xlsx",
        "fields": [
            {"header": "权限", "required": "否", "description": "可留空，系统默认按导师导入。若填写，固定写 TEACHER。", "example": "TEACHER"},
            {"header": "姓名", "required": "是", "description": "导师姓名。", "example": "张晓农"},
            {"header": "电子邮箱", "required": "否", "description": "可选，推荐填写学校邮箱。", "example": "zhangxn@mail.sysu.edu.cn"},
            {"header": "工号", "required": "是", "description": "导师工号。若系统中已存在，将覆盖更新原导师信息，不会新建重复账号。", "example": "190001"},
            {"header": "职称", "required": "是", "description": "必填，例如教授、副教授。", "example": "教授"},
            {"header": "研究方向", "required": "否", "description": "可填写主要研究方向，支持较长文本。", "example": "作物遗传育种、农业信息化"},
        ],
    },
    "student": {
        "sheet_title": "学生导入模板",
        "filename": "学生信息导入模板.xlsx",
        "fields": [
            {"header": "权限", "required": "否", "description": "可留空，系统默认按学生导入。若填写，固定写 STUDENT。", "example": "STUDENT"},
            {"header": "姓名", "required": "是", "description": "学生姓名。", "example": "李明轩"},
            {"header": "电子邮箱", "required": "否", "description": "可选，推荐填写学校邮箱。", "example": "limx58@mail.sysu.edu.cn"},
            {"header": "学号", "required": "是", "description": "学生学号。若系统中已存在，将覆盖更新原学生信息，不会新建重复账号。", "example": "24330001"},
            {"header": "年级", "required": "是", "description": "必填，填写培养年级，例如 2024。", "example": "2024"},
        ],
    },
}

HEADER_FILL = PatternFill("solid", fgColor="0B6B3A")
EXAMPLE_FILL = PatternFill("solid", fgColor="EAF5EE")
NOTE_FILL = PatternFill("solid", fgColor="F5F8FA")


def build_workbook(kind: str) -> Workbook:
    config = TEMPLATES[kind]
    workbook = Workbook()

    template_sheet = workbook.active
    template_sheet.title = config["sheet_title"]

    headers = [field["header"] for field in config["fields"]]
    examples = [field["example"] for field in config["fields"]]

    template_sheet.append(headers)
    template_sheet.append(examples)

    for column_index, field in enumerate(config["fields"], start=1):
        header_cell = template_sheet.cell(row=1, column=column_index)
        header_cell.font = Font(color="FFFFFF", bold=True)
        header_cell.fill = HEADER_FILL
        header_cell.alignment = Alignment(horizontal="center", vertical="center")

        example_cell = template_sheet.cell(row=2, column=column_index)
        example_cell.fill = EXAMPLE_FILL
        example_cell.alignment = Alignment(vertical="center", wrap_text=True)

        column_width = max(len(field["header"]) * 2.2, len(field["example"]) * 1.2, len(field["description"]) * 0.55, 16)
        template_sheet.column_dimensions[get_column_letter(column_index)].width = min(column_width, 42)

    template_sheet.freeze_panes = "A2"
    template_sheet.row_dimensions[1].height = 24
    template_sheet.row_dimensions[2].height = 26

    guide_sheet = workbook.create_sheet("字段说明")
    guide_sheet.append(["字段", "是否必填", "说明", "示例"])
    for field in config["fields"]:
        guide_sheet.append([field["header"], field["required"], field["description"], field["example"]])

    for column_index in range(1, 5):
        cell = guide_sheet.cell(row=1, column=column_index)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in guide_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[2].fill = NOTE_FILL

    guide_sheet.column_dimensions["A"].width = 18
    guide_sheet.column_dimensions["B"].width = 12
    guide_sheet.column_dimensions["C"].width = 58
    guide_sheet.column_dimensions["D"].width = 28
    guide_sheet.freeze_panes = "A2"

    return workbook


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: generate_import_template.py <teacher|student> <output_path>", file=sys.stderr)
        return 1

    kind = sys.argv[1].strip().lower()
    output_path = Path(sys.argv[2]).expanduser().resolve()

    if kind not in TEMPLATES:
        print(f"Unsupported template type: {kind}", file=sys.stderr)
        return 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(kind)
    workbook.save(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
